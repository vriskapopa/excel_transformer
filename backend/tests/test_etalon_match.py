"""Compare generated 18233 export totals against эталон ready files."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.services.export_18233_templates import (
    export_18233_from_templates,
    fill_specification_template,
)
from app.services.materials_18233 import kit_templates
from app.services.profile_18233 import parse_bundle, reconcile_18233
from app.services.reconcile import items_to_dicts
from app.services.materials_18233 import kit_sources


def _spec_totals(path: Path) -> dict[str, float]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for r in range(20, ws.max_row + 1):
        val = str(ws.cell(r, 2).value or "")
        if "total" in val.lower():
            return {
                "rolls": float(ws.cell(r, 3).value or 0),
                "meters": float(ws.cell(r, 8).value or 0),
                "area": float(ws.cell(r, 10).value or 0),
                "nw": float(ws.cell(r, 11).value or 0),
                "gw": float(ws.cell(r, 12).value or 0),
                "amount": float(ws.cell(r, 14).value or 0),
            }
    raise AssertionError(f"No total row in {path}")


def _articles(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    arts = []
    for r in range(23, ws.max_row + 1):
        val = ws.cell(r, 2).value
        art = ws.cell(r, 4).value
        if val is not None and "total" in str(val).lower():
            break
        if art:
            arts.append(str(art).strip())
    return arts


def _run_kit(kit: str, expected_count: int, tmp_path: Path) -> None:
    src = kit_sources(kit)
    bundle = parse_bundle([src["invoice"], src["packing"], src["specification"]])
    items = items_to_dicts(reconcile_18233(bundle))
    product_items = [
        i for i in items if (i.get("source_traces") or {}).get("invoice") or (i.get("source_traces") or {}).get("specification")
    ]
    assert len(product_items) == expected_count

    out_dir = tmp_path / kit
    header = {"invoice_no": f"ZFRMB26148-{kit}", "invoice_date": "Aug.19,2026"}
    paths = export_18233_from_templates(
        items,
        out_dir,
        header=header,
        shipment_title="18233",
    )
    assert len(paths) == 2
    assert all("СПЕЦИФ" not in p.name.upper() for p in paths)
    assert any("ИНВОЙС" in p.name.upper() or "инвойс" in p.name.lower() for p in paths)
    assert any("ПАКИНГ" in p.name.upper() or "пакинг" in p.name.lower() for p in paths)

    # Spec with colors is not exported; still check numbers vs эталон via filler
    generated_spec = fill_specification_template(
        kit_templates(kit)["specification"],
        out_dir / f"_check_spec_{kit}.xlsx",
        items,
        header,
        kit,
    )
    etalon_spec = kit_templates(kit)["specification"]

    got_arts = _articles(generated_spec)
    exp_arts = _articles(etalon_spec)
    assert got_arts == exp_arts, (got_arts, exp_arts)

    got = _spec_totals(generated_spec)
    exp = _spec_totals(etalon_spec)
    for key in ("rolls", "meters", "nw", "gw", "amount"):
        assert abs(got[key] - exp[key]) <= 0.15, f"{kit} {key}: got={got[key]} exp={exp[key]}"
    assert abs(got["area"] - exp["area"]) <= 0.5, f"{kit} area: {got['area']} vs {exp['area']}"


def test_6261_matches_etalon(tmp_path: Path) -> None:
    _run_kit("626-1", 13, tmp_path)


def test_6262_matches_etalon(tmp_path: Path) -> None:
    _run_kit("626-2", 17, tmp_path)
