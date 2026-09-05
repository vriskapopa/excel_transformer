from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.models.enums import DocType, ErrorSeverity, ErrorType
from app.parsing.schemas import ParsedDocument, ParsedLine
from app.services.catalog import CatalogIndex
from app.services.reconcile import reconcile_documents


def _doc(filename: str, doc_type: DocType, lines: list[dict]) -> ParsedDocument:
    return ParsedDocument(
        filename=filename,
        file_path=filename,
        doc_type=doc_type,
        lines=[ParsedLine.model_validate(line) for line in lines],
        text_preview=filename,
    )


def test_beijing_catalog_fill_md812(tmp_path: Path) -> None:
    from openpyxl import Workbook

    catalog_path = tmp_path / "(описание )сводная.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Article", "TN VED", "Description EN", "Description RU"])
    ws.append(
        [
            "MD 812",
            "7318230009",
            "Furniture metal rivet 8 x 12 mm",
            "Заклепка мебельная ступенчатая 8 × 12 мм",
        ]
    )
    wb.save(catalog_path)
    catalog = CatalogIndex.from_excel(catalog_path)

    invoice = _doc(
        "Invoice.xlsx",
        DocType.INVOICE,
        [
            {
                "row_index": 1,
                "article": "MD 812",
                "normalized_article": "MD812",
                "raw": {
                    "Article": "MD 812",
                    "Qty": 200100,
                    "Unit Price": 0.0592,
                    "Amount": 11845.92,
                },
            }
        ],
    )
    items = reconcile_documents([invoice], catalog=catalog)
    assert len(items) == 1
    assert items[0].customs_data["tnved_code"] == "7318230009"
    assert "Furniture metal rivet" in (items[0].customs_data.get("description_en") or "")


def test_catalog_missing_is_yellow_not_invented() -> None:
    invoice = _doc(
        "Invoice.xlsx",
        DocType.INVOICE,
        [
            {
                "row_index": 1,
                "article": "MD 811",
                "normalized_article": "MD811",
                "raw": {"Article": "MD 811", "Qty": 10, "Price": 1},
            }
        ],
    )
    catalog = CatalogIndex()
    items = reconcile_documents([invoice], catalog=catalog)
    assert items[0].customs_data.get("tnved_code") in (None, "")
    assert any(f.error_type == ErrorType.CATALOG_NOT_FOUND for f in items[0].flags)
    assert any(f.severity == ErrorSeverity.YELLOW for f in items[0].flags)


def test_hs_vs_catalog_mismatch_is_red(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "catalog.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Article", "TN VED", "Description EN"])
    ws.append(["Noble 110", "5407613000", "Fabric"])
    wb.save(path)
    catalog = CatalogIndex.from_excel(path)

    invoice = _doc(
        "626-1-INVOICE.xlsx",
        DocType.INVOICE,
        [
            {
                "row_index": 1,
                "article": "Noble 110",
                "normalized_article": "NOBLE110",
                "raw": {
                    "Article": "Noble 110",
                    "HS Code": "5407610000",
                    "Meters": 611.5,
                    "Rolls": 15,
                },
            }
        ],
    )
    items = reconcile_documents([invoice], catalog=catalog)
    assert any(
        f.error_type == ErrorType.MISMATCH and f.severity == ErrorSeverity.RED for f in items[0].flags
    )


def test_row_without_article_not_merged() -> None:
    invoice = _doc(
        "inv.xlsx",
        DocType.INVOICE,
        [
            {
                "row_index": 1,
                "article": "A1",
                "normalized_article": "A1",
                "raw": {"Article": "A1", "Qty": 1},
            },
            {
                "row_index": 2,
                "article": None,
                "normalized_article": "",
                "raw": {"Qty": 5, "Price": 10},
            },
        ],
    )
    items = reconcile_documents([invoice])
    assert [i.normalized_article for i in items] == ["A1"]
    assert all(i.normalized_article for i in items)


def test_packing_aggregate_distribution() -> None:
    invoice_110 = _doc(
        "626-1-INVOICE.xlsx",
        DocType.INVOICE,
        [
            {
                "row_index": 1,
                "article": "Noble 110",
                "normalized_article": "NOBLE110",
                "raw": {"Article": "Noble 110", "Meters": 611.5, "Rolls": 15},
            },
            {
                "row_index": 2,
                "article": "Noble 624",
                "normalized_article": "NOBLE624",
                "raw": {"Article": "Noble 624", "Meters": 507.0, "Rolls": 12},
            },
        ],
    )
    packing = _doc(
        "626-1-PL.xlsx",
        DocType.PACKING_LIST,
        [
            {
                "row_index": 1,
                "article": "Noble",
                "normalized_article": "NOBLE",
                "raw": {
                    "Article": "Noble",
                    "Rolls": 27,
                    "Meters": 1118.5,
                    "Net Weight": 1062.58,
                    "Gross Weight": 1092,
                },
            }
        ],
    )
    items = reconcile_documents([invoice_110, packing])
    by_key = {i.normalized_article: i for i in items}
    child = by_key["NOBLE110"]
    assert child.packing_data.get("net_weight") is not None
    assert child.packing_data.get("distributed_from_article") == "NOBLE"
    # Share ≈ 611.5 / 1118.5
    expected = round(1062.58 * (611.5 / 1118.5), 2)
    assert child.packing_data["net_weight"] == expected
