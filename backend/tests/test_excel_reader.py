from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.parsing.excel_reader import read_excel
from app.parsing.pipeline import parse_file


def test_read_excel_normalizes_article(tmp_path: Path) -> None:
    path = tmp_path / "invoice.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Invoice"
    sheet.append(["Article", "Qty", "Price"])
    sheet.append(["MO SHO-01", 10, 12.5])
    sheet.append(["mo  sho-01", 2, 12.5])
    book.save(path)

    sheets, lines, preview = read_excel(str(path))
    assert sheets == ["Invoice"]
    assert "Invoice" in preview
    assert [line["normalized_article"] for line in lines] == ["MOSHO-01", "MOSHO-01"]
    assert lines[0]["article"] == "MO SHO-01"


def test_parse_file_classifies_invoice_excel(tmp_path: Path) -> None:
    path = tmp_path / "Commercial_Invoice.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Invoice"
    sheet.append(["Article", "Quantity", "Unit Price", "Amount"])
    sheet.append(["ABC-9", 1, 10, 10])
    book.save(path)

    document = parse_file(str(path), allow_ocr=False)
    assert document.doc_type is not None
    assert document.doc_type.value == "INVOICE"
    assert document.lines[0].normalized_article == "ABC-9"


def test_parse_nameless_excel_keeps_prices(tmp_path: Path) -> None:
    from app.services.reconcile import items_to_dicts, reconcile_documents

    path = tmp_path / "goods.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["MD 812", 200, 0.0592, 11.84])
    sheet.append(["AB-90", 10, 12.5, 125])
    book.save(path)

    document = parse_file(str(path), allow_ocr=False)
    assert document.doc_type.value == "INVOICE"
    items = items_to_dicts(reconcile_documents([document]))
    by_article = {item["article"]: item for item in items}
    assert by_article["MD 812"]["commercial_data"]["price"] == 0.0592
    assert by_article["AB-90"]["commercial_data"]["qty"] == 10


def test_parse_file_reads_every_sheet(tmp_path: Path) -> None:
    path = tmp_path / "two_sheets.xlsx"
    book = Workbook()
    first = book.active
    first.title = "cover"
    first.append(["Supplier", "Tosun"])
    second = book.create_sheet("спец")
    second.append(["Articul/Артикул", "Qty", "Price"])
    second.append(["ZIMMY 162", 10, 5.78])
    third = book.create_sheet("описание")
    third.append(["Articul/Артикул", "Qty", "Price"])
    third.append(["SINDRI 162", 4, 6.17])
    book.save(path)

    document = parse_file(str(path), allow_ocr=False)
    articles = {line.article for line in document.lines}
    assert "ZIMMY 162" in articles
    assert "SINDRI 162" in articles


