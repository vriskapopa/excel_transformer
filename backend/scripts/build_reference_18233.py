"""Build article reference (customs code + bilingual description) from эталон specs."""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT_BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT_BACKEND))

from openpyxl import load_workbook

from app.parsing.normalize import normalize_article

ROOT = Path(__file__).resolve().parents[2] / "materials" / "MVP_18233"
OUT = ROOT_BACKEND / "app" / "services" / "data" / "reference_18233.json"


def kit_ready_dir(kit: str) -> Path:
    base = next(p for p in ROOT.rglob("*") if p.is_dir() and p.name.endswith(kit))
    return next(c for c in base.iterdir() if c.is_dir() and c.name.startswith("02"))


def extract_from_spec(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]  # инвойс 1
    # find header row
    header_row = None
    for r in range(1, 40):
        vals = [str(ws.cell(r, c).value or "").lower() for c in range(1, 15)]
        if any("art" in v for v in vals) and any("customs" in v or "таможен" in v for v in vals):
            header_row = r
            break
    if header_row is None:
        raise RuntimeError(f"No header in {path}")

    result: dict[str, dict] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        article = ws.cell(r, 4).value
        if article is None:
            continue
        text = str(article).strip()
        if not text or text.lower().startswith("total"):
            break
        key = normalize_article(text)
        result[key] = {
            "article": text,
            "description": ws.cell(r, 5).value,
            "hs_code": str(ws.cell(r, 6).value or "").replace(".0", ""),
            "tnved_code": str(ws.cell(r, 7).value or "").replace(".0", ""),
            "country": ws.cell(r, 16).value,
            "manufacturer": ws.cell(r, 17).value,
        }
        # normalize numeric-looking codes
        for field in ("hs_code", "tnved_code"):
            val = result[key][field]
            if val.endswith(".0"):
                result[key][field] = val[:-2]
            if isinstance(ws.cell(r, 6 if field == "hs_code" else 7).value, (int, float)):
                result[key][field] = str(int(ws.cell(r, 6 if field == "hs_code" else 7).value))
    return result


def main() -> None:
    catalog: dict[str, dict] = {}
    for kit in ("626-1", "626-2"):
        ready = kit_ready_dir(kit)
        spec = next(p for p in ready.iterdir() if "СПЕЦИФ" in p.name.upper() or "специф" in p.name.lower())
        part = extract_from_spec(spec)
        catalog.update(part)
        print(kit, len(part), spec.name)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", OUT, "articles", len(catalog))


if __name__ == "__main__":
    main()
