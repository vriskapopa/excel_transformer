"""Create minimal demo Excel files matching Word-spec examples."""

from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1] / "sample_data"


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)

    inv = Workbook()
    ws = inv.active
    ws.title = "Invoice"
    ws.append(["Article", "Color", "Qty", "Unit", "Unit Price", "Amount", "HS Code"])
    ws.append(["MD 812", "silver", 200100, "pcs", 0.0592, 11845.92, ""])
    ws.append(["MO SHO-01", "black", 50, "pcs", 1.2, 60, ""])
    ws.append(["MD 811", "black", 10, "pcs", 0.5, 5, ""])  # missing in catalog → yellow
    inv.save(ROOT / "Invoice n Packing list.xlsx")

    pl = Workbook()
    # same file second sheet style — separate packing file for clarity
    ws = pl.active
    ws.title = "Packing list"
    ws.append(["Article", "Boxes", "Net Weight", "Gross Weight", "Volume"])
    ws.append(["MD 812", 20, 120.5, 130.0, 1.2])
    ws.append(["MOSHO-01", 2, 5.0, 5.5, 0.1])
    pl.save(ROOT / "Packing list.xlsx")

    catalog = Workbook()
    ws = catalog.active
    ws.title = "Catalog"
    ws.append(["Article", "TN VED", "Description EN", "Description RU"])
    ws.append(
        [
            "MD 812",
            "7318230009",
            "Furniture metal rivet 8 x 12 mm",
            "Заклепка мебельная ступенчатая 8 × 12 мм",
        ]
    )
    ws.append(
        [
            "MOSHO-01",
            "8302420000",
            "Furniture fitting",
            "Фурнитура мебельная",
        ]
    )
    catalog.save(ROOT / "(описание )сводная.xlsx")

    # 18233-style mini kit
    inv182 = Workbook()
    ws = inv182.active
    ws.title = "Invoice"
    ws.append(["Article", "Rolls", "Meters", "Area", "Unit Price", "Amount", "HS Code"])
    ws.append(["Noble 110", 15, 611.5, 868.33, 43.8, 26783.70, "5407610000"])
    ws.append(["Noble 624", 12, 507.0, 720.0, 43.8, 22206.60, "5407610000"])
    inv182.save(ROOT / "626-1-YS-RMB-EXW-INVOICE.xlsx")

    pl182 = Workbook()
    ws = pl182.active
    ws.title = "Packing List"
    ws.append(["Article", "Rolls", "Meters", "Net Weight", "Gross Weight"])
    ws.append(["Noble", 27, 1118.5, 1062.58, 1092])
    pl182.save(ROOT / "626-1-YS-RMB-EXW-PL.xlsx")

    cat182 = Workbook()
    ws = cat182.active
    ws.append(["Article", "TN VED", "Description EN", "Description RU"])
    ws.append(["Noble 110", "5407613000", "Woven fabric", "Ткань"])
    ws.append(["Noble 624", "5407613000", "Woven fabric", "Ткань"])
    cat182.save(ROOT / "18233_catalog.xlsx")

    print(f"Sample files written to {ROOT}")


if __name__ == "__main__":
    main()
