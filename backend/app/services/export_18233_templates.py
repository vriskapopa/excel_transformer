"""Export 18233 by filling copies of customer эталон templates."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.parsing.normalize import normalize_article
from app.services.materials_18233 import kit_templates, materials_available
from app.services.profile_18233 import _design_family


def _r2(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return round(float(value), 2)


def _r3(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return round(float(value), 3)


def _detect_kit(items: list[dict[str, Any]], header: dict[str, Any] | None) -> str:
    for item in items:
        sub = item.get("invoice_subkit") or (item.get("commercial_data") or {}).get("invoice_subkit")
        if sub in {"626-1", "626-2"}:
            return str(sub)
    inv = (header or {}).get("invoice_no") or ""
    match = re.search(r"(626-\d)", str(inv))
    if match:
        return match.group(1)
    # fallback by article set
    arts = {normalize_article(i.get("article") or "") for i in items}
    if any(a.startswith("NAPPA") or a.startswith("MAGIC") for a in arts):
        return "626-2"
    return "626-1"


def _ordered_products(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    products = [
        i
        for i in items
        if (i.get("source_traces") or {}).get("invoice") or (i.get("source_traces") or {}).get("specification")
    ]
    products.sort(key=lambda i: int(((i.get("source_traces") or {}).get("invoice") or {}).get("no") or 999))
    return products


def _apply_header(ws, header: dict[str, Any], kit: str) -> None:
    if not header:
        return
    invoice_no = header.get("invoice_no") or f"ZFRMB26148-{kit}"
    date = header.get("invoice_date") or header.get("date") or ""
    contract = header.get("contract_no") or ""
    # Best-effort replacements in known cells / scanned labels
    for row in ws.iter_rows(min_row=1, max_row=25, max_col=15):
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            if "инв номер" in text.lower() or text.strip().upper() in {"INV.NO.", "INV.NO"}:
                right = ws.cell(cell.row, cell.column + 1)
                if right.value is not None or "инв" in text.lower():
                    right.value = invoice_no
            if text.strip().upper() == "DATE:" and date:
                ws.cell(cell.row, cell.column + 1).value = date
            if contract and text.lower().startswith("contract") and "№" in text or (
                contract and "контракт" in text.lower()
            ):
                # keep structure, optionally rewrite if operator provided full line
                if header.get("contract_line"):
                    cell.value = header["contract_line"]


def _set_cell(ws, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        # write into the top-left of the merge that owns this cell
        for merged in ws.merged_cells.ranges:
            if cell.coordinate in merged:
                ws.cell(merged.min_row, merged.min_col).value = value
                return
        return
    cell.value = value


def _clear_block(ws, start_row: int, end_row: int, min_col: int, max_col: int) -> None:
    # Unmerge any ranges intersecting the block so we can rewrite freely
    to_unmerge = []
    for merged in ws.merged_cells.ranges:
        if merged.max_row < start_row or merged.min_row > end_row:
            continue
        if merged.max_col < min_col or merged.min_col > max_col:
            continue
        to_unmerge.append(str(merged))
    for ref in to_unmerge:
        ws.unmerge_cells(ref)
    for r in range(start_row, end_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def fill_specification_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb[wb.sheetnames[0]]
    products = _ordered_products(items)

    # data starts at row 23 (from dump), header at 22
    data_start = 23
    # clear old product rows until Total
    total_row = None
    for r in range(data_start, ws.max_row + 1):
        val = ws.cell(r, 2).value
        if val is not None and "total" in str(val).lower():
            total_row = r
            break
    clear_end = (total_row - 1) if total_row else data_start + max(30, len(products))
    _clear_block(ws, data_start, clear_end, 2, 20)

    sums = {"rolls": 0.0, "meters": 0.0, "area": 0.0, "nw": 0.0, "gw": 0.0, "amount": 0.0}
    manufacturer = (header or {}).get("manufacturer") or "HANGZHOU ZHONGFANG TEXTILE IMP/EXP.CO.,LTD"
    country = (header or {}).get("country") or "КИТАЙ"

    for idx, item in enumerate(products, start=1):
        row = data_start + idx - 1
        c = item.get("commercial_data") or {}
        p = item.get("packing_data") or {}
        u = item.get("customs_data") or {}
        article = item.get("article") or ""
        rolls = _r2(p.get("rolls")) or 0
        meters = _r2(p.get("meters")) or 0
        width = p.get("width")
        area = _r3(p.get("area"))
        if area is None and meters and width:
            area = _r3(float(meters) * float(width))
        nw = _r2(p.get("net_weight"))
        gw = _r2(p.get("gross_weight"))
        price = c.get("price")
        amount = c.get("amount")
        if amount is None and price is not None and meters:
            amount = _r2(float(price) * float(meters))
        hs = u.get("hs_code") or ""
        tnved = u.get("tnved_code") or hs
        desc = u.get("description_ru") or u.get("description_en") or u.get("description") or ""
        if u.get("description_en") and u.get("description_ru") and not desc:
            desc = f"{u['description_en']}/{u['description_ru']}"
        if u.get("description"):
            desc = u["description"]

        ws.cell(row, 2).value = idx
        ws.cell(row, 3).value = rolls
        ws.cell(row, 4).value = article
        ws.cell(row, 5).value = desc
        ws.cell(row, 6).value = hs
        ws.cell(row, 7).value = tnved
        ws.cell(row, 8).value = meters
        ws.cell(row, 9).value = width
        ws.cell(row, 10).value = area
        ws.cell(row, 11).value = nw
        ws.cell(row, 12).value = gw
        ws.cell(row, 13).value = price
        ws.cell(row, 14).value = amount
        ws.cell(row, 15).value = tnved
        ws.cell(row, 16).value = u.get("country") or country
        ws.cell(row, 17).value = u.get("manufacturer") or manufacturer
        ws.cell(row, 18).value = desc
        ws.cell(row, 19).value = article
        ws.cell(row, 20).value = hs

        sums["rolls"] += float(rolls or 0)
        sums["meters"] += float(meters or 0)
        sums["area"] += float(area or 0)
        sums["nw"] += float(nw or 0)
        sums["gw"] += float(gw or 0)
        sums["amount"] += float(amount or 0)

    # write / update total row
    if total_row is None:
        total_row = data_start + len(products)
    ws.cell(total_row, 2).value = "Total/Итого:"
    ws.cell(total_row, 3).value = _r2(sums["rolls"])
    ws.cell(total_row, 8).value = _r2(sums["meters"])
    ws.cell(total_row, 10).value = _r3(sums["area"])
    ws.cell(total_row, 11).value = _r2(sums["nw"])
    ws.cell(total_row, 12).value = _r2(sums["gw"])
    ws.cell(total_row, 14).value = _r2(sums["amount"])

    # title / invoice no
    inv = (header or {}).get("invoice_no") or f"ZFRMB26148-{kit}"
    date = (header or {}).get("invoice_date") or (header or {}).get("date") or "Aug.19,2026"
    for r in range(1, 15):
        for c in range(1, 15):
            val = ws.cell(r, c).value
            if val is None:
                continue
            text = str(val)
            if text.startswith("SPECIFICATION №") or text.startswith("SPECIFICATION"):
                ws.cell(r, c).value = f"SPECIFICATION №{inv} dd {date}"
            if "инв номер" in text.lower():
                ws.cell(r, c + 1).value = inv
            if text.strip().lower() in {"дата:", "date:"}:
                ws.cell(r, c + 1).value = date

    # also refresh description sheet if present
    if len(wb.sheetnames) > 1:
        desc_ws = wb[wb.sheetnames[1]]
        d_start = 13
        for r in range(1, 20):
            if str(desc_ws.cell(r, 4).value or "").lower().startswith("art"):
                d_start = r + 1
                break
        _clear_block(desc_ws, d_start, d_start + 40, 2, 15)
        for idx, item in enumerate(products, start=1):
            row = d_start + idx - 1
            c = item.get("commercial_data") or {}
            p = item.get("packing_data") or {}
            u = item.get("customs_data") or {}
            desc = u.get("description") or u.get("description_en") or u.get("description_ru") or ""
            _set_cell(desc_ws, row, 2, idx)
            _set_cell(desc_ws, row, 3, _r2(p.get("rolls")))
            _set_cell(desc_ws, row, 4, item.get("article"))
            _set_cell(desc_ws, row, 5, desc)
            _set_cell(desc_ws, row, 6, u.get("hs_code"))
            _set_cell(desc_ws, row, 7, u.get("tnved_code") or u.get("hs_code"))
            _set_cell(desc_ws, row, 8, _r2(p.get("meters")))
            _set_cell(desc_ws, row, 9, p.get("width"))
            _set_cell(desc_ws, row, 10, _r3(p.get("area")))
            _set_cell(desc_ws, row, 11, _r2(p.get("net_weight")))
            _set_cell(desc_ws, row, 12, _r2(p.get("gross_weight")))
            _set_cell(desc_ws, row, 13, c.get("price"))
            _set_cell(desc_ws, row, 14, c.get("amount"))

    wb.save(output)
    return output


def fill_invoice_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
    packing_groups: list[dict[str, Any]] | None = None,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb.active
    products = _ordered_products(items)

    # Find table header
    header_row = None
    for r in range(1, 40):
        if str(ws.cell(r, 1).value or "").strip().upper() == "NO." and str(ws.cell(r, 2).value or "").upper() == "DESIGN":
            header_row = r
            break
    if header_row is None:
        header_row = 24
    data_start = header_row + 1

    # clear until TOTAL / empty tail
    end = data_start
    for r in range(data_start, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a is None and b is None:
            # allow a few blanks then stop
            if r > data_start + 5:
                break
        end = r
        if b is not None and str(b).upper().startswith("TOTAL"):
            end = r - 1
            break
    _clear_block(ws, data_start, max(end, data_start + 40), 1, 10)

    # Build family groups in invoice appearance order
    families: list[str] = []
    by_family: dict[str, list[dict[str, Any]]] = {}
    for item in products:
        fam = _design_family(item.get("article") or "")
        if fam not in by_family:
            families.append(fam)
            by_family[fam] = []
        by_family[fam].append(item)

    row = data_start
    no = 1
    for fam in families:
        group_items = by_family[fam]
        group_rolls = 0.0
        group_meters = 0.0
        group_area = 0.0
        price = None
        hs = None
        width = None
        for item in group_items:
            inv = (item.get("source_traces") or {}).get("invoice") or {}
            p = item.get("packing_data") or {}
            c = item.get("commercial_data") or {}
            u = item.get("customs_data") or {}
            ws.cell(row, 1).value = no
            ws.cell(row, 2).value = item.get("article")
            ws.cell(row, 3).value = u.get("hs_code") or inv.get("hs_code")
            ws.cell(row, 4).value = _r2(p.get("rolls") or inv.get("rolls"))
            ws.cell(row, 5).value = p.get("width") or inv.get("width")
            ws.cell(row, 6).value = _r3(p.get("area") or inv.get("area"))
            ws.cell(row, 7).value = _r2(p.get("meters") or inv.get("meters"))
            ws.cell(row, 8).value = c.get("unit") or inv.get("unit") or "meters"
            # product lines in эталон leave price/amount empty
            ws.cell(row, 9).value = None
            ws.cell(row, 10).value = None
            group_rolls += float(p.get("rolls") or inv.get("rolls") or 0)
            group_meters += float(p.get("meters") or inv.get("meters") or 0)
            group_area += float(p.get("area") or inv.get("area") or 0)
            price = c.get("price") or inv.get("price") or price
            hs = u.get("hs_code") or inv.get("hs_code") or hs
            width = p.get("width") or inv.get("width") or width
            no += 1
            row += 1

        # group summary row
        prefix = "ARTIFICIAL LEATHER" if fam.startswith("NAPPA") or fam.startswith("MAGIC") else "SOFA FABRIC"
        # Magic/Nappa naming from эталон packing
        label_name = group_items[0].get("article", "").split()[0]
        ws.cell(row, 1).value = None
        ws.cell(row, 2).value = f"{prefix}\n{label_name}"
        ws.cell(row, 3).value = hs
        ws.cell(row, 4).value = _r2(group_rolls)
        ws.cell(row, 5).value = width
        ws.cell(row, 6).value = _r3(group_area)
        ws.cell(row, 7).value = _r2(group_meters)
        ws.cell(row, 8).value = "meters"
        ws.cell(row, 9).value = price
        ws.cell(row, 10).value = _r2(float(price) * group_meters) if price is not None else None
        row += 1

    # TOTAL
    ws.cell(row, 2).value = "TOTAL:"
    ws.cell(row, 4).value = _r2(sum(float((i.get("packing_data") or {}).get("rolls") or 0) for i in products))
    ws.cell(row, 6).value = _r3(sum(float((i.get("packing_data") or {}).get("area") or 0) for i in products))
    ws.cell(row, 7).value = _r2(sum(float((i.get("packing_data") or {}).get("meters") or 0) for i in products))
    ws.cell(row, 10).value = _r2(
        sum(float((i.get("commercial_data") or {}).get("amount") or 0) for i in products)
    )

    inv_no = (header or {}).get("invoice_no") or f"ZFRMB26148-{kit}"
    date = (header or {}).get("invoice_date") or (header or {}).get("date") or "Aug.19,2026"
    for r in range(1, header_row):
        for c in range(1, 12):
            val = ws.cell(r, c).value
            if val is None:
                continue
            if str(val).strip().upper() in {"INV.NO.", "INV.NO"}:
                ws.cell(r, c + 1).value = inv_no
            if str(val).strip().upper() == "DATE:":
                ws.cell(r, c + 1).value = date
            if str(val).startswith("SPECIFICATION №"):
                ws.cell(r, c).value = f"SPECIFICATION № {inv_no} {date}"

    wb.save(output)
    return output


def fill_packing_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb.active
    products = _ordered_products(items)

    header_row = None
    for r in range(1, 20):
        if str(ws.cell(r, 1).value or "").strip().upper() in {"NO.", "NO"} and "DESIGN" in str(
            ws.cell(r, 2).value or ""
        ).upper():
            header_row = r
            break
    if header_row is None:
        header_row = 9
    data_start = header_row + 1

    # clear
    for r in range(data_start, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").upper() == "TOTAL":
            break
        for c in range(1, 9):
            ws.cell(r, c).value = None

    # Prefer original Packing List group lines (incl. split Magic rows)
    groups: list[dict[str, Any]] = []
    seen = set()
    for item in products:
        pl = (item.get("source_traces") or {}).get("packing_list_group") or {}
        parts = pl.get("parts") or [pl]
        for part in parts:
            if not part:
                continue
            sig = (
                part.get("row_index"),
                part.get("design"),
                part.get("rolls"),
                part.get("meters"),
            )
            if sig in seen:
                continue
            seen.add(sig)
            groups.append(part)

    if not groups:
        # fallback aggregate by family
        families: list[str] = []
        by_family: dict[str, list[dict[str, Any]]] = {}
        for item in products:
            fam = _design_family(item.get("article") or "")
            if fam not in by_family:
                families.append(fam)
                by_family[fam] = []
            by_family[fam].append(item)
        for fam in families:
            group_items = by_family[fam]
            groups.append(
                {
                    "design": f"SOFA FABRIC\n{(group_items[0].get('article') or fam).split()[0]}",
                    "rolls": sum(float((i.get("packing_data") or {}).get("rolls") or 0) for i in group_items),
                    "meters": sum(float((i.get("packing_data") or {}).get("meters") or 0) for i in group_items),
                    "net_weight": sum(
                        float((i.get("packing_data") or {}).get("net_weight") or 0) for i in group_items
                    ),
                    "gross_weight": sum(
                        float((i.get("packing_data") or {}).get("gross_weight") or 0) for i in group_items
                    ),
                    "area": sum(float((i.get("packing_data") or {}).get("area") or 0) for i in group_items),
                }
            )

    row = data_start
    sums = {"rolls": 0.0, "meters": 0.0, "nw": 0.0, "gw": 0.0, "area": 0.0}
    for idx, pl in enumerate(groups, start=1):
        rolls = pl.get("rolls")
        meters = pl.get("meters")
        nw = pl.get("net_weight")
        gw = pl.get("gross_weight")
        area = pl.get("area")
        gm = pl.get("gm")
        design = pl.get("design")
        ws.cell(row, 1).value = idx
        ws.cell(row, 2).value = design
        ws.cell(row, 3).value = gm
        ws.cell(row, 4).value = _r2(rolls)
        ws.cell(row, 5).value = _r2(meters)
        ws.cell(row, 6).value = _r2(nw)
        ws.cell(row, 7).value = _r2(gw)
        ws.cell(row, 8).value = _r3(area)
        sums["rolls"] += float(rolls or 0)
        sums["meters"] += float(meters or 0)
        sums["nw"] += float(nw or 0)
        sums["gw"] += float(gw or 0)
        sums["area"] += float(area or 0)
        row += 1

    ws.cell(row, 1).value = "TOTAL"
    ws.cell(row, 4).value = _r2(sums["rolls"])
    ws.cell(row, 5).value = _r2(sums["meters"])
    ws.cell(row, 6).value = _r2(sums["nw"])
    ws.cell(row, 7).value = _r2(sums["gw"])
    ws.cell(row, 8).value = _r3(sums["area"])

    inv_no = (header or {}).get("invoice_no") or f"ZFRMB26148-{kit}"
    date = (header or {}).get("invoice_date") or (header or {}).get("date") or "Aug.19,2026"
    for r in range(1, header_row):
        for c in range(1, 9):
            val = ws.cell(r, c).value
            if val is None:
                continue
            if str(val).strip().upper() in {"INV.NO.", "INV.NO"}:
                ws.cell(r, c + 1).value = inv_no
            if str(val).strip().upper() == "DATE:":
                ws.cell(r, c + 1).value = date

    wb.save(output)
    return output


def export_18233_from_templates(
    items: list[dict[str, Any]],
    output_dir: Path,
    header: dict[str, Any] | None = None,
    shipment_title: str = "18233",
) -> list[Path]:
    if not materials_available():
        raise FileNotFoundError("MVP_18233 materials not found; cannot use эталон templates")

    kit = _detect_kit(items, header)
    templates = kit_templates(kit)
    output_dir.mkdir(parents=True, exist_ok=True)

    inv_out = output_dir / f"{shipment_title} ИНВОЙС {kit}.xlsx"
    pl_out = output_dir / f"{shipment_title} ПАКИНГ {kit}.xlsx"
    spec_out = output_dir / f"{shipment_title} {kit} СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx"

    paths = [
        fill_invoice_template(templates["invoice"], inv_out, items, header, kit),
        fill_packing_template(templates["packing"], pl_out, items, header, kit),
        fill_specification_template(templates["specification"], spec_out, items, header, kit),
    ]
    return paths
